"""Build the editable S4 report from frozen experiment logs (python-docx).

This authoring utility does not train models or change scientific results.
Run from any directory with an artifact-authoring Python that has python-docx.
"""
import json
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'report'
URL = 'https://github.com/huyryan220806/random-forest-optuna-optimization'
base = json.loads((ROOT / 'results/baseline.json').read_text(encoding='utf-8'))
best = json.loads((ROOT / 'results/best_config.json').read_text(encoding='utf-8'))
final = json.loads((ROOT / 'results/final_metrics.json').read_text(encoding='utf-8'))
b, o = final['models']['baseline'], final['models']['optimized']
assert final['status'] == best['status'] == 'complete'
assert best['best_params'] == {'n_estimators': 50, 'max_depth': 16, 'min_samples_leaf': 1}
f6 = lambda value: f'{value:.6f}'

doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.top_margin, s.bottom_margin = Cm(2), Cm(1.8)
s.left_margin = s.right_margin = Cm(2.2)
s.header_distance = s.footer_distance = Cm(0.9)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption']:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.15
normal.paragraph_format.space_after = Pt(7)
doc.styles['Title'].font.size = Pt(25)
doc.styles['Title'].font.bold = True
doc.styles['Heading 1'].font.size = Pt(18)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(12)
doc.styles['Heading 2'].font.size = Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(10)
doc.styles['Caption'].font.size = Pt(9)
doc.styles['Caption'].font.italic = True
doc.styles['Caption'].font.bold = False
# The bundled default template carries a Title paragraph border. Remove it.
for border in list(doc.styles.element.xpath('.//w:pBdr')):
    border.getparent().remove(border)
doc.core_properties.title = 'Tối ưu siêu tham số Random Forest bằng Optuna'
doc.core_properties.author = '(tự điền)'
doc.core_properties.subject = 'Bài tập lớn cá nhân Nhập môn Tối ưu hóa'
s.different_first_page_header_footer = True
hp = s.header.paragraphs[0]
hp.text = 'NHẬP MÔN TỐI ƯU HÓA     •     RANDOM FOREST VÀ OPTUNA'
hp.runs[0].font.size = Pt(8)
fp = s.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fp.add_run('Trang ').font.size = Pt(9)
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); fp._p.append(fld)

def p(text, style=None):
    return doc.add_paragraph(text, style)

def h(text):
    doc.add_heading(text, level=2)

def page(title):
    doc.add_page_break()
    doc.add_heading(title, level=1)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Cm(width)
    for i, text in enumerate(headers):
        t.rows[0].cells[i].text = str(text)
    for row in rows:
        for c, text in zip(t.add_row().cells, row):
            c.text = str(text)
    for ri, row in enumerate(t.rows):
        pr = row._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit'); pr.append(no_split)
        if ri == 0:
            pr.append(OxmlElement('w:tblHeader'))
        for ci, cell in enumerate(row.cells):
            cell.width = Cm(widths[ci])
            tcpr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:fill'), '17365D' if ri == 0 else ('F3F5F7' if ri % 2 else 'FFFFFF'))
            tcpr.append(shd)
            margins = OxmlElement('w:tcMar')
            for edge in ['top', 'left', 'bottom', 'right']:
                el = OxmlElement('w:' + edge); el.set(qn('w:w'), '80'); el.set(qn('w:type'), 'dxa'); margins.append(el)
            tcpr.append(margins)
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right']:
                el = OxmlElement('w:' + edge)
                for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                    el.set(qn('w:' + key), value)
                borders.append(el)
            tcpr.append(borders)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.05
                for run in para.runs:
                    run.font.size = Pt(10)
                    run.font.bold = ri == 0
                    run.font.color.rgb = RGBColor.from_string('FFFFFF' if ri == 0 else '000000')
    p('') .paragraph_format.space_after = Pt(0)
    return t

def picture(name, caption):
    pp = doc.add_paragraph()
    pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pp.add_run().add_picture(str(ROOT / 'figures' / name), width=Cm(16.4))
    p(caption, 'Caption')

def code(text):
    pp = p(text)
    pp.paragraph_format.space_after = Pt(8)
    pp.paragraph_format.line_spacing = 1.0
    for run in pp.runs:
        run.font.name = 'Consolas'; run.font.size = Pt(9)

def math_run(text):
    r = OxmlElement('m:r'); t = OxmlElement('m:t'); t.text = text; r.append(t); return r

def objective():
    # Native OMML fraction, summation and subscript; checked in the PDF render.
    pp = p(''); pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    eq = OxmlElement('m:oMath')
    lim = OxmlElement('m:limLow')
    e = OxmlElement('m:e'); e.append(math_run('min')); lim.append(e)
    l = OxmlElement('m:lim'); l.append(math_run('x ∈ Ω')); lim.append(l); eq.append(lim)
    eq.append(math_run(' f(x) = −'))
    frac = OxmlElement('m:f')
    for tag, val in [('num', '1'), ('den', '3')]:
        el = OxmlElement('m:' + tag); el.append(math_run(val)); frac.append(el)
    eq.append(frac)
    nary = OxmlElement('m:nary')
    pr = OxmlElement('m:naryPr'); char = OxmlElement('m:chr'); char.set(qn('m:val'), '∑'); pr.append(char); nary.append(pr)
    for tag, val in [('sub', 'k = 1'), ('sup', '3')]:
        el = OxmlElement('m:' + tag); el.append(math_run(val)); nary.append(el)
    e = OxmlElement('m:e'); sub = OxmlElement('m:sSub')
    se = OxmlElement('m:e'); se.append(math_run('AUC')); sub.append(se)
    si = OxmlElement('m:sub'); si.append(math_run('k')); sub.append(si)
    e.append(sub); e.append(math_run('(x)')); nary.append(e); eq.append(nary); pp._p.append(eq)

# 1 Cover and abstract
p('TRƯỜNG ĐẠI HỌC VĂN LANG\nKHOA CÔNG NGHỆ THÔNG TIN').alignment = WD_ALIGN_PARAGRAPH.CENTER
p('BÀI TẬP LỚN CÁ NHÂN\nNHẬP MÔN TỐI ƯU HÓA', 'Subtitle').alignment = WD_ALIGN_PARAGRAPH.CENTER
p('Tối ưu siêu tham số\nRandom Forest bằng Optuna\ncho bài toán phân loại', 'Title')
p('Hyperparameter Optimization of Random Forest Using Optuna for Classification', 'Subtitle')
p('Họ và tên: (tự điền)\nMSSV: (tự điền)\nLớp: (tự điền)\nGiảng viên hướng dẫn: (tự điền)')
p('Bản bàn giao S4 • 30/10/2026\nThực nghiệm chính: 25/10/2026')
h('Tóm tắt')
p(f'Đề tài khảo sát liệu Optuna TPE với 30 trials có cải thiện Random Forest mặc định trên Breast Cancer Wisconsin Diagnostic hay không. '
  f'Cùng tập train, cùng 3 folds và seed 42 được dùng cho hai cấu hình. Best CV ROC-AUC đạt {f6(o["cv_auc_mean"])}, '
  f'so với {f6(b["cv_auc_mean"])} của baseline. Cấu hình chọn từ CV có 50 cây, độ sâu tối đa 16 và lá tối thiểu 1 mẫu. '
  'Holdout cho AUC tăng nhẹ nhưng F1 và confusion matrix không đổi. Báo cáo phân tích cả lợi ích nhỏ, chi phí tìm kiếm và giới hạn của một seed.')
h('Kho mã nguồn và kết quả')
p(URL)
p('Mã nguồn, CSV/JSON, hai hình, tài liệu và lệnh chạy lại được lưu cùng repository. Thông tin cá nhân còn để trống theo yêu cầu của người thực hiện.', 'Caption')

# 2 Problem and formulation
page('1 Bối cảnh và mô hình hóa')
h('Câu hỏi nghiên cứu')
p('Cấu hình mặc định là điểm xuất phát hợp lý nhưng chưa chắc phù hợp nhất với một tập dữ liệu cụ thể. Bài tập đặt câu hỏi: trong ngân sách nhỏ và quy trình đánh giá cố định, TPE có tìm được Random Forest cho ROC-AUC cross-validation cao hơn mặc định không, và lợi ích đó đổi lấy chi phí nào?')
p('Trọng tâm là quy trình tối ưu có thể kiểm chứng: biến quyết định, miền khả thi, hàm mục tiêu, baseline, lịch sử thử nghiệm và quyết định cuối. MVP chỉ dùng một dataset, một họ mô hình và một phương pháp HPO; không bổ sung giao diện, mô hình khác hay nhiều seed vào kết quả chính.')
h('Biến quyết định và miền khả thi')
p('Vector x gồm ba siêu tham số n_estimators, max_depth và min_samples_leaf. Miền Ω là tích Descartes của các tập giá trị dưới đây; có 5 × 7 × 8 = 280 cấu hình khả thi.')
table(['Biến', 'Miền tìm kiếm', 'Baseline'], [
    ['n_estimators', '50, 100, 150, 200, 250', '100'],
    ['max_depth', 'None, 4, 6, 8, 12, 16, 24', 'None'],
    ['min_samples_leaf', 'Số nguyên từ 1 đến 8', '1'],
], [5.0, 8.0, 3.6])
h('Hàm mục tiêu')
objective()
p('AUCₖ(x) là ROC-AUC của cấu hình x trên validation fold k, chỉ lấy từ tập train. Code dùng direction="maximize" cho trung bình AUC, tương đương cực tiểu hóa giá trị âm ở công thức. F1 và accuracy là chỉ tiêu báo cáo holdout, không tham gia objective.')
h('Ngân sách và chi phí')
p('Giới hạn cứng là x thuộc Ω và tối đa 30 trials trong run chính. Mỗi trial thực hiện 3 lần fit; baseline cần thêm 3 lần fit ngoài budget search. Trial lặp cấu hình vẫn tiêu thụ budget. Thời gian search, fit và predict được đo để phân tích đánh đổi, không có ngưỡng latency để loại nghiệm.')
p('Baseline nằm trong miền tìm kiếm nhưng không được enqueue vào study. Vì vậy đây là so sánh cấu hình mặc định với HPO có chi phí bổ sung, không phải so sánh hai thuật toán tìm kiếm có cùng budget.')

# 3 Dataset and protocol
page('2 Dữ liệu và protocol thực nghiệm')
p('Breast Cancer Wisconsin Diagnostic là benchmark công khai từ UCI [1], được nạp bằng load_breast_cancer của scikit-learn [2]. Dữ liệu gồm 569 mẫu và 30 đặc trưng số mô tả đặc điểm nhân tế bào từ ảnh chọc hút kim nhỏ. UCI công bố giấy phép CC BY 4.0; repository ghi nguồn và cách nạp trong data/README.md.')
p('Bản dùng trong thí nghiệm là dữ liệu đóng gói cùng scikit-learn 1.7.2. Code không tải dữ liệu ngoài lúc chạy, không chọn đặc trưng và không chuẩn hóa. Kiểm tra các giá trị đầu vào hữu hạn và kích thước đúng được thực hiện trong tests. Đây là bài tập phương pháp, không xây dựng công cụ chẩn đoán.')
h('Quy ước nhãn')
p('scikit-learn gốc dùng malignant=0, benign=1. Dự án đổi thành y = (target == 0).astype(int), do đó malignant là lớp dương 1, benign là 0. AUC lấy xác suất lớp 1; F1 và mọi diễn giải confusion matrix dùng cùng quy ước này.')
table(['Tập', 'Benign 0', 'Malignant 1', 'Tổng'], [
    ['Train', 285, 170, 455], ['Holdout test', 72, 42, 114], ['Toàn bộ', 357, 212, 569]
], [5.4, 3.6, 4.0, 3.6])
h('Phân chia và bảo vệ holdout')
p('train_test_split tách một lần theo tỉ lệ 80/20, stratify=y, shuffle=True, random_state=42. Trên 455 mẫu train, StratifiedKFold dùng 3 folds, shuffle=True và seed 42. Hai fold validation có 152 mẫu, fold còn lại có 151 mẫu. Danh sách chỉ số folds được tạo một lần rồi dùng chung cho baseline và mọi trial.')
p('Objective chỉ nhận X_train, y_train, danh sách folds và cấu hình. Holdout không được truyền vào closure tối ưu. Sau khi kết thúc 30 trials, cấu hình thắng được xác định từ CV, ghi vào checkpoint, rồi mới fit trên toàn bộ train và dự đoán test.')
p('Seed cho split, CV, Random Forest và TPE đều bằng 42. Tất cả chạy tuần tự với n_jobs=1. Cố định seed giúp kiểm chứng trên cùng phiên bản; không thay thế việc đo độ ổn định qua nhiều seed độc lập.')
h('Nguyên tắc so sánh')
p('Không điều chỉnh search space, budget, seed hay ngưỡng sau khi xem holdout. Chạy lại trong thư mục riêng chỉ nhằm kiểm tra tái lập. Dữ liệu test không được dùng để chọn lại nghiệm; khi CV hòa, quy tắc chọn giữ baseline đã được cố định trước đánh giá.')

# 4 Method
page('3 Random Forest và phương pháp TPE')
h('Baseline Random Forest')
p('Random Forest kết hợp nhiều cây quyết định học từ các mẫu bootstrap và chọn ngẫu nhiên tập đặc trưng xét tại nút. Trong scikit-learn, xác suất phân lớp của rừng được lấy trung bình từ xác suất các cây [3]. Sự ngẫu nhiên làm các cây khác nhau, còn phép tổng hợp giúp giảm phụ thuộc vào một cây đơn lẻ.')
p('Baseline dùng RandomForestClassifier mặc định: 100 cây, max_depth=None và min_samples_leaf=1; các tham số khác giữ mặc định, ngoại trừ random_state=42 và n_jobs=1. Tăng số cây làm tăng lượng tính toán. Giới hạn độ sâu hoặc tăng số mẫu tối thiểu ở lá có thể giảm độ phức tạp nhưng cũng có thể làm mất cấu trúc hữu ích trong dữ liệu.')
h('TPE trong Optuna')
p('Tree-structured Parzen Estimator sử dụng lịch sử thử nghiệm để đề xuất tham số [4, 5]. Thay vì chỉ lấy ngẫu nhiên mọi cấu hình, TPE tách các quan sát thành nhóm có objective tốt và nhóm còn lại, ước lượng mật độ tham số l(x) và g(x), rồi ưu tiên ứng viên có tỷ lệ l(x)/g(x) cao. Với maximize, nhóm tốt tương ứng AUC cao.')
p('Dự án dùng framework Optuna [6], phiên bản 4.5.0, TPESampler(seed=42, n_startup_trials=10). Mười trials đầu là khởi tạo ngẫu nhiên; 20 trials còn lại khai thác lịch sử theo TPE. Thiết lập multivariate giữ mặc định False. Không pruning, không timeout và không enqueue baseline; NopPruner bảo đảm trial được đánh giá đầy đủ nếu không phát sinh lỗi.')
h('Luồng tính toán')
code('1. Nạp dữ liệu, đổi nhãn, tách train/test.\n2. Tạo 3 folds trên train; đánh giá baseline.\n3. Lặp 30 trials: đề xuất x thuộc Ω, tạo RF mới,\n   tính AUC từng fold, trả về mean, ghi CSV.\n4. Chọn trial có CV mean lớn nhất; lưu best_config.\n5. Ghi lựa chọn theo CV, fit hai model trên train,\n   dự đoán holdout một lần/model; lưu final_metrics.\n6. Đọc log để sinh hình và bảng so sánh.')
h('Điều kiện dừng và lỗi')
p('Run chính dừng sau 30 trials thành công. Nếu có lỗi, CSV ghi trạng thái FAIL và thông tin lỗi, không gán điểm giả; chương trình dừng với mã lỗi và best_config đánh dấu incomplete. Run incomplete không được dùng để đánh giá holdout. Smoke 3 trials chỉ kiểm tra đường chạy và chưa kiểm chứng giai đoạn TPE sau khởi tạo.')
p('Chọn TPE vì nó hỗ trợ trực tiếp biến nguyên và biến phân loại, phù hợp ngân sách nhỏ và dễ lưu lịch sử. Không có đối chứng Random Search cùng budget trong MVP, nên kết quả này không chứng minh TPE vượt trội hơn các phương pháp HPO khác.')

# 5 Implementation
page('4 Triển khai và lưu dấu thực nghiệm')
p('Hai module giữ quy trình đơn giản: src/experiment.py chứa cấu hình, split, baseline, objective, logging và đánh giá cuối; src/plots.py đọc log để dựng hai hình và bảng. ExperimentConfig và SEARCH_SPACE là nguồn cấu hình của code, tránh phân tán các giá trị seed và miền tham số.')
table(['Tệp', 'Thông tin được lưu'], [
    ['results/baseline.json', 'Tham số mặc định, AUC từng fold/mean/std, timing, seed, split/CV và môi trường.'],
    ['results/trials.csv', '30 dòng trial với params, AUC, std, điểm folds, duration, state, seed và error.'],
    ['results/best_config.json', 'Winner theo CV, toàn bộ tham số model, sampler, budget và chi phí search.'],
    ['results/final_metrics.json', 'Lựa chọn trước test, SHA-256 nguồn S2, metric holdout, confusion matrix, fit/predict timing.'],
    ['figures/', 'Hai PNG và comparison.csv/md sinh bằng code từ log.']
], [6.1, 10.5])
h('Ghi checkpoint và tính nhất quán')
p('CSV được cập nhật sau mỗi trial. Các JSON/CSV được ghi qua tệp tạm rồi thay thế để giảm nguy cơ đọc tệp đang ghi dở. Trong CSV, max_depth=None được lưu là chuỗi None; JSON dùng null. Khi nạp log, code kiểm tra trạng thái hoàn tất, số trials, metadata, điểm fold, cấu hình mặc định và winner theo CV.')
p('Trước đánh giá holdout, dấu SHA-256 của ba log S2 được lưu cùng lựa chọn model. Phạm vi băm là văn bản UTF-8 chuẩn hóa CRLF thành LF để không sai khác khi checkout trên Windows. Checkpoint thêm từng model sau khi đánh giá; chạy lại --evaluate-only trên kết quả complete dùng dữ liệu đã lưu, không fit hoặc predict lại.')
h('Môi trường')
p('Run chính dùng Python 3.12.10 trên Windows 11 AMD64, 12 logical CPUs; RF và CV đều chỉ dùng một worker. Phiên bản chính: NumPy 2.2.6, SciPy 1.15.3, scikit-learn 1.7.2, Optuna 4.5.0, matplotlib 3.10.6. requirements.txt ghim dependencies trực tiếp; timing đo bằng time.perf_counter().')
h('Kiểm chứng mã thực tế')
p('Bộ 24 tests bao phủ nhãn/split/folds, giữ mặc định baseline, đường CV chung, logging và trial lỗi, tính metric, checkpoint holdout, kiểm tra nguồn và sinh hình. Tests S3 dùng dữ liệu tổng hợp trong thư mục tạm; các số đó không được đưa vào bảng chính. Chi tiết lần chạy sạch và phạm vi tái lập được ghi ở report/REPRODUCIBILITY.md.')

# 6 CV results
page('5 Kết quả cross validation')
table(['Đại lượng', 'Baseline', 'Best Optuna'], [
    ['CV AUC mean', f6(b['cv_auc_mean']), f6(o['cv_auc_mean'])],
    ['CV AUC std', f6(b['cv_auc_std']), f6(o['cv_auc_std'])],
    ['Số cây / độ sâu / lá', '100 / None / 1', '50 / 16 / 1'],
    ['Số trial search', '0', str(best['n_trials'])],
], [7.0, 4.8, 4.8])
picture('best_so_far.png', 'Hình 1. Best-so-far CV ROC-AUC theo số trial hoàn thành. Nguồn: results/trials.csv và baseline.json; sinh bằng src.plots.')
p(f'Best trial là #{best["best_trial_number"]} khi đếm từ 0, tức lần đánh giá thứ {best["best_trial_number"] + 1}. '
  f'Chênh lệch mean so với baseline là {f6(best["cv_auc_improvement"])} AUC, tương đương khoảng {best["cv_auc_improvement"] * 100:.4f} điểm phần trăm. '
  'Các trials tiếp theo không tạo được giá trị cao hơn trong ngân sách 30 lần thử.')
p('Đường best-so-far luôn không giảm theo định nghĩa; đoạn nằm ngang chỉ cho biết chưa tìm được nghiệm tốt hơn trong các lần thử còn lại, không chứng minh đã đạt cực trị toàn cục của 280 cấu hình.')
p('Std giữa ba folds dùng ddof=0. Optimized có mean cao hơn nhưng std cũng cao hơn baseline. Đây là độ phân tán giữa các folds của cùng một split, không phải khoảng tin cậy hay độ ổn định của nhiều lần chạy độc lập.')

# 7 Costs
page('6 Chi phí tính toán')
picture('auc_vs_runtime.png', 'Hình 2. AUC CV và thời gian từng trial. Ngôi sao là trial đã chọn; đường baseline chỉ tham chiếu AUC, không có tọa độ thời gian baseline.')
table(['Phép đo', 'Baseline', 'Optimized'], [
    ['Search tăng thêm (giây)', '0', f6(final['search_duration_sec'])],
    ['Fit trên 455 mẫu (giây)', f6(b['fit_duration_sec']), f6(o['fit_duration_sec'])],
    ['Predict batch 114 mẫu (ms)', f'{b["inference_duration_sec"] * 1000:.3f}', f'{o["inference_duration_sec"] * 1000:.3f}'],
], [8.0, 4.3, 4.3])
p('Hình 2 lấy duration_sec của trial, gồm sampling, tạo model, CV và ghi metric vào trial. Thời gian CV riêng được lưu ở cv_duration_sec. Search_duration_sec bao gồm toàn bộ study.optimize và checkpoint CSV, nhưng không gồm baseline CV. Vì phạm vi đo khác nhau, không đặt baseline thành một điểm thời gian trên scatter.')
p(f'Baseline CV của run chính mất {f6(base["duration_sec"])} giây. Optuna cần thêm {final["search_duration_sec"]:.3f} giây search để có mức tăng CV nhỏ. '
  'Cấu hình 50 cây trong run này có fit và predict nhanh hơn baseline 100 cây. Số cây là một chỉ dấu độ phức tạp; chưa đo kích thước model hoặc bộ nhớ.')
p('Mỗi timing fit/predict chỉ đo một lần và chịu ảnh hưởng tải máy, cache cùng phiên bản môi trường. Predict là một batch 114 mẫu, không phải latency một yêu cầu riêng lẻ. Không dùng các số đo này như ràng buộc cứng, bảo đảm tốc độ hay kết luận thống kê về hiệu năng.')

# 8 Holdout
page('7 Đánh giá trên holdout')
p('Cấu hình optimized đã được chốt bởi CV trước khi dự đoán test. Mỗi model fit trên toàn bộ 455 mẫu train rồi gọi predict_proba một lần cho 114 mẫu holdout. AUC dùng xác suất malignant; nhãn phân loại dùng ngưỡng cố định 0.5, bằng ngưỡng thì chọn benign theo argmax của scikit-learn.')
table(['Chỉ số test', 'Baseline', 'Optimized'], [
    ['ROC-AUC', f6(b['test_roc_auc']), f6(o['test_roc_auc'])],
    ['F1 malignant', f6(b['test_f1']), f6(o['test_f1'])],
    ['Accuracy', f6(b['test_accuracy']), f6(o['test_accuracy'])],
    ['Balanced accuracy', f6(b['test_balanced_accuracy']), f6(o['test_balanced_accuracy'])],
], [7.0, 4.8, 4.8])
p('Bảng lấy từ results/final_metrics.json; bảng máy đọc được ở figures/comparison.csv.', 'Caption')
h('Confusion matrix của cả hai model')
table(['Nhãn thật', 'Dự đoán benign 0', 'Dự đoán malignant 1'], [
    ['Benign 0', b['confusion_matrix'][0][0], b['confusion_matrix'][0][1]],
    ['Malignant 1', b['confusion_matrix'][1][0], b['confusion_matrix'][1][1]],
], [5.4, 5.6, 5.6])
p('Hàng là nhãn thật, cột là dự đoán. Cả hai có TN=72, FP=0, FN=3, TP=39. Từ đó accuracy=111/114 và F1=78/81. Mỗi cấu hình bỏ sót 3 mẫu malignant; ma trận tổng hợp không cho biết đó có phải cùng các mẫu hay không. AUC cao không có nghĩa không bỏ sót lớp dương.')
h('Diễn giải sự khác nhau giữa AUC và F1')
p(f'Chênh lệch test AUC là {o["test_roc_auc"] - b["test_roc_auc"]:.6f}. AUC đánh giá thứ tự của xác suất qua nhiều ngưỡng, còn F1 được tính từ nhãn tại ngưỡng đã chọn. '
  'Do đó thứ tự hoặc mức xác suất có thể thay đổi mà F1 tại 0.5 vẫn không đổi. Kết quả này chỉ cho thấy cải thiện nhẹ về xếp hạng trên split đang xét.')
p('Holdout có 42 mẫu malignant và 72 benign nên biến động ở một số ít mẫu có thể làm thay đổi metric. Không dùng kết quả test để chỉnh ngưỡng nhằm làm bảng đẹp hơn, không lựa chọn seed hoặc trial theo test và không tuyên bố khác biệt có ý nghĩa thống kê.')

# 9 Decision and limits
page('8 Quyết định và giới hạn')
h('Nghiệm cuối cho phạm vi dự án')
p('Giữ cấu hình Optuna với n_estimators=50, max_depth=16, min_samples_leaf=1. Lựa chọn tuân theo quy tắc CV đã cố định trước khi xem test: CV mean cao hơn thì chọn optimized, nếu hòa giữ baseline. Cấu hình này có ít cây hơn, và trong run đã ghi nhận, thời gian fit/predict cũng thấp hơn.')
p('Holdout đóng vai trò đánh giá sau lựa chọn. AUC test nhỉnh hơn nhưng F1, accuracy, balanced accuracy và số mẫu bỏ sót không cải thiện. Vì vậy kết luận là TPE tìm được cấu hình có điểm CV cao hơn một lượng nhỏ trong budget đã đặt; chưa có bằng chứng về lợi ích phân loại rõ rệt hay ưu thế ổn định trên dữ liệu mới.')
h('Chi phí của mức cải thiện')
p('Baseline đã mạnh và không tiêu tốn search. Để tăng khoảng 0.000396 CV AUC, HPO mất thêm gần 25 giây trên máy chạy chính. Đây là chi phí nhỏ cho bài tập nhưng vẫn phải tính vào so sánh. Việc model cuối nhanh hơn không xóa đi chi phí tìm kiếm đã phát sinh; không suy rộng lợi ích sang hệ thống lớn khi chưa đo.')
h('Các giới hạn của thiết kế')
p('Thứ nhất, chỉ có một dataset nhỏ và một seed. Các folds cùng thuộc một tập train nên std ba folds không thể thay cho phân tích nhiều seed. Best CV còn chịu ảnh hưởng của việc chọn điểm lớn nhất qua 30 lần thử, vì thế có thể lạc quan hơn đánh giá độc lập.')
p('Thứ hai, budget 30 nhỏ hơn 280 cấu hình khả thi và có thể chứa cấu hình lặp. Đường hội tụ không chứng minh cực trị toàn cục. Không có Random Search cùng budget hoặc kiểm định giả thuyết; không thể kết luận TPE là thuật toán HPO tốt nhất từ thí nghiệm này.')
p('Thứ ba, thời gian được đo một lần, không đánh giá bộ nhớ, kích thước model, calibration hay độ bền trước thay đổi phân phối. Dữ liệu là benchmark công khai, không đại diện cho quy trình bệnh viện cụ thể và không đủ để xác nhận hiệu quả chẩn đoán lâm sàng.')
h('Giá trị bàn giao')
p('Giá trị của MVP là một chuỗi có thể kiểm tra từ công thức đến code, từ CSV/JSON đến hai hình và quyết định. Người đọc có thể kiểm tra baseline có bị làm bất lợi hay không, test có bị dùng để tuning hay không và chi phí search được tính ở đâu. Những phần mở rộng như nhiều seed hoặc Random Search chỉ thực hiện khi có yêu cầu tiếp theo.')

# 10 Reproduction and AI
page('9 Chạy lại và trách nhiệm kiểm chứng')
h('Chuẩn bị môi trường')
p('Clone repository bằng URL ở bìa, mở PowerShell tại thư mục gốc. Dùng Python 3.12 để sát môi trường đã kiểm chứng. Dữ liệu có sẵn trong scikit-learn; chỉ cần mạng khi clone và cài dependencies. Không cần GPU hoặc thay đổi ExecutionPolicy.')
code('py -3.12 -m venv .venv\n.\\.venv\\Scripts\\python.exe -m pip install -r requirements.txt\n.\\.venv\\Scripts\\python.exe -m pip check\n.\\.venv\\Scripts\\python.exe -m pytest -q')
h('Kiểm tra nhanh và xem kết quả đã đóng băng')
code('.\\.venv\\Scripts\\python.exe -m src.experiment --trials 3 --seed 42 --output-dir results/smoke\n.\\.venv\\Scripts\\python.exe -m src.experiment --evaluate-only --input-dir results\n.\\.venv\\Scripts\\python.exe -m src.plots --input-dir results --output-dir figures')
p('Lệnh evaluate-only trên log complete đọc lại kết quả; không đánh giá test thêm. Sinh hình chỉ dùng log. Smoke được lưu riêng và không dùng thay kết quả chính. Bản demo 5–7 phút được hướng dẫn chi tiết trong report/DEMO.md.')
h('Tái lập đầy đủ bằng thư mục mới')
code('.\\.venv\\Scripts\\python.exe -m src.experiment --trials 30 --seed 42 --output-dir results/repro/lan-moi\n.\\.venv\\Scripts\\python.exe -m src.experiment --evaluate-only --input-dir results/repro/lan-moi\n.\\.venv\\Scripts\\python.exe -m src.plots --input-dir results/repro/lan-moi --output-dir figures/repro/lan-moi')
p('Mỗi lần chạy đầy đủ cần tên thư mục mới nếu thư mục cũ đã có final_metrics.json. So sánh params, seed, trạng thái, điểm từng fold/CV và metric holdout; không đòi timing hoặc timestamp giống nhau. Lần chạy lại chỉ là kiểm chứng cấu hình cố định, không mở vòng tuning trên test.')
h('AI assistance và phần sinh viên tự thực hiện')
p('Codex hỗ trợ thiết kế cấu trúc, viết code baseline/TPE/evaluation/plots, chạy kiểm tra và soạn tài liệu S1–S4. Các kiểm chứng tự động được lưu trong nhật ký và checklist tái lập. Không xem đó là bằng chứng người học đã tự hiểu hoặc tự chạy toàn bộ bài.')
p('Phần sinh viên tự kiểm chứng: (tự điền). Trước khi nộp, người thực hiện cần điền thông tin bìa, tự chạy lệnh, đối chiếu bảng/hình, giải thích objective và TPE, luyện demo và nộp link lên LMS. Lịch commit thật không thay thế yêu cầu checkpoint 5 tuần của giảng viên; không chỉnh ngày commit để giả tiến độ.')

# 11 References
page('10 Tài liệu tham khảo')
refs = [
    ('[1] Wolberg, W., Mangasarian, O., Street, N., và Street, W. (1993). Breast Cancer Wisconsin (Diagnostic). UCI Machine Learning Repository. DOI 10.24432/C5DW2B. Giấy phép CC BY 4.0.', 'https://doi.org/10.24432/C5DW2B'),
    ('[2] scikit-learn developers. load_breast_cancer. Tài liệu phiên bản 1.7.2. Nguồn mô tả bộ dữ liệu, kích thước và nhãn.', 'https://scikit-learn.org/1.7/modules/generated/sklearn.datasets.load_breast_cancer.html'),
    ('[3] scikit-learn developers. RandomForestClassifier. Tài liệu phiên bản 1.7.2. Nguồn tham số mặc định và cơ chế predict_proba.', 'https://scikit-learn.org/1.7/modules/generated/sklearn.ensemble.RandomForestClassifier.html'),
    ('[4] Bergstra, J., Bardenet, R., Bengio, Y., và Kégl, B. (2011). Algorithms for Hyper-Parameter Optimization. Advances in Neural Information Processing Systems 24.', 'https://papers.nips.cc/paper/4443-algorithms-for-hyper-parameter-optimization.pdf'),
    ('[5] Optuna contributors. optuna.samplers.TPESampler. Tài liệu phiên bản 4.5.0. Nguồn nguyên lý TPE, seed và n_startup_trials.', 'https://optuna.readthedocs.io/en/v4.5.0/reference/samplers/generated/optuna.samplers.TPESampler.html'),
    ('[6] Akiba, T., Sano, S., Yanase, T., Ohta, T., và Koyama, M. (2019). Optuna: A Next-generation Hyperparameter Optimization Framework. KDD 2019.', 'https://arxiv.org/abs/1907.10902'),
]
for desc, url in refs:
    p(desc)
    pp = p(url)
    for run in pp.runs: run.font.size = Pt(9)
h('Nguồn số liệu và ngày truy cập')
p('Các nguồn trực tuyến được đối chiếu ngày 30/10/2026. Mọi metric và timing trong báo cáo lấy từ thực nghiệm chính ngày 25/10/2026: results/baseline.json, trials.csv, best_config.json và final_metrics.json. Hai hình được sinh bằng src/plots.py. Số liệu không lấy từ bài báo hay bảng benchmark bên ngoài.')
p('Repository: ' + URL)
p('Tài liệu yêu cầu học phần được cung cấp cục bộ trong plans/. PROJECT_PLAN.md ở gốc là đặc tả và nhật ký làm việc; báo cáo này bám phạm vi MVP đã chốt. Không tự nhận đã hoàn thành các thủ tục nộp bài hoặc kiểm chứng cá nhân thay cho sinh viên.')

OUT.mkdir(exist_ok=True)
path = OUT / 'Bao_cao_Random_Forest_Optuna.docx'
doc.save(path)
print(path)
